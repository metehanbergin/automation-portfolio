"""Text PDF, XLSX, CSV and MIME-email extraction with deterministic accounting rules."""
import csv
import hashlib
import io
import re
import time
from decimal import Decimal, InvalidOperation
from email import policy
from email.parser import BytesParser

from openpyxl import load_workbook
from pypdf import PdfReader

from .core import audit, now, notify, uid


def initial():
    return dict(customers=[{'id':'C101','name':'Northline Studio'}, {'id':'C102','name':'Cedar Works'},
                           {'id':'C103','name':'Harbor Design'}],
                invoices=[dict(id=f'INV-{1001+i}', customer=['C101','C102','C103'][i%3],
                               total=amount, paid=0, currency='USD')
                          for i, amount in enumerate([125000,84000,210000,64000,95000,180000,72000,32000])],
                documents=[], postings=[])


def money(value):
    try:
        amount = Decimal(str(value).replace(',', '').replace('$','').strip())
        if not amount.is_finite() or amount <= 0 or amount.as_tuple().exponent < -2:
            raise ValueError('Amount must be positive with at most two decimal places')
        return int(amount * 100)
    except (InvalidOperation, TypeError):
        raise ValueError('Invalid payment amount')


def parse_attachment(name, raw):
    ext = name.lower().rsplit('.',1)[-1]
    if ext == 'pdf':
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 20:
            raise ValueError('PDF exceeds 20-page demo limit')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        if not text.strip():
            raise ValueError('No text layer. Scanned PDFs require an OCR provider.')
        return text
    if ext == 'xlsx':
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        if ws.max_row and ws.max_row > 1000:
            raise ValueError('Spreadsheet exceeds 1,000-row demo limit')
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if len(rows) < 2:
            raise ValueError('Spreadsheet needs a header and a payment row')
        if len(rows) > 2:
            raise ValueError('One payment per upload; split a batch into individual payment rows')
        return '\n'.join(f'{k}: {v}' for k,v in zip(rows[0],rows[1]))
    if ext == 'csv':
        rows = list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
        if len(rows) != 1:
            raise ValueError('CSV must contain exactly one payment row')
        return '\n'.join(f'{k}: {v}' for k,v in rows[0].items())
    if ext == 'eml':
        msg = BytesParser(policy=policy.default).parsebytes(raw)
        pieces = []
        for part in msg.walk():
            if part.get_filename():
                pieces.append(parse_attachment(part.get_filename(), part.get_payload(decode=True)))
            elif part.get_content_type() == 'text/plain':
                pieces.append(part.get_content())
        if not pieces:
            raise ValueError('Email has no supported text or attachment')
        return '\n'.join(pieces)
    if ext in ('txt','text'):
        return raw.decode('utf-8')
    raise ValueError('Unsupported attachment. Use text PDF, XLSX, CSV, EML or TXT.')


def extract(text):
    fields = {}
    for key in ('customer','invoice','amount','reference','currency'):
        matches = re.findall(rf'(?im)^[ \t]*{key}[ \t]*:[ \t]*([^\r\n]*)',text)
        unique = set(x.strip() for x in matches if x.strip())
        if len(unique) > 1:
            raise ValueError(f'Conflicting {key} fields; review the source')
        fields[key] = next(iter(unique), '')
    if not fields['amount']:
        raise ValueError('No payment amount found')
    fields['amount'] = money(fields['amount'])
    fields['currency'] = fields['currency'].upper()
    return fields


def process(state, name, raw):
    started = time.perf_counter()
    book = state['remittance']
    doc = dict(id=uid('REM'), name=name, hash=hashlib.sha256(raw).hexdigest(), at=now(),
               status='review', reasons=[], fields={}, confidence=0, source='', resolution=None)
    book['documents'].append(doc)
    audit(state,'remittance',doc['id'],'Received',f'{name} · {len(raw)} bytes')
    if any(x['hash']==doc['hash'] for x in book['documents'][:-1]):
        doc['status']='duplicate'; doc['reasons']=['Duplicate document fingerprint']
    else:
        try:
            doc['source'] = parse_attachment(name,raw)
            doc['fields'] = extract(doc['source'])
            f=doc['fields']
            audit(state,'remittance',doc['id'],'Extracted','Customer, invoice, amount, reference and currency parsed')
            customers=[c for c in book['customers'] if f['customer'].casefold() in (c['id'].casefold(),c['name'].casefold())]
            inv=next((i for i in book['invoices'] if i['id']==f['invoice']),None)
            if not f['reference']: doc['reasons'].append('Missing payment reference')
            if len(customers)!=1: doc['reasons'].append('Unmatched customer')
            if not inv: doc['reasons'].append('Missing or unknown invoice')
            if f['currency']!='USD': doc['reasons'].append('Currency missing or unsupported')
            if inv and customers and inv['customer']!=customers[0]['id']: doc['reasons'].append('Invoice belongs to a different customer')
            if inv and f['amount']!=inv['total']-inv['paid']:
                doc['reasons'].append('Partial payment' if f['amount']<inv['total']-inv['paid'] else 'Amount exceeds outstanding balance')
            duplicate = any(p['reference']==f['reference'] for p in book['postings']) if f['reference'] else False
            if duplicate:
                doc['status']='duplicate'; doc['reasons']=['Payment reference already posted']
            doc['confidence']=max(0,100-len(doc['reasons'])*22)
            if not doc['reasons']:
                post(state,doc,inv,f['reference'],f['amount'],'system')
        except Exception as exc:
            doc['status']='review';doc['reasons']=[f'Malformed or unsupported attachment: {str(exc)[:180]}']
    doc['duration_ms']=round((time.perf_counter()-started)*1000,2)
    audit(state,'remittance',doc['id'],'Decision',f"{doc['status']}: {', '.join(doc['reasons']) or 'All deterministic checks passed'}")
    return doc


def post(state,doc,inv,reference,amount,actor):
    book=state['remittance']
    if any(p['reference']==reference for p in book['postings']):
        raise ValueError('Payment reference already posted')
    if amount<=0 or amount>inv['total']-inv['paid']:
        raise ValueError('Payment exceeds remaining invoice balance')
    inv['paid']+=amount
    book['postings'].append(dict(id=uid('PAY'),document=doc['id'],invoice=inv['id'],amount=amount,reference=reference,at=now()))
    doc['status']='posted' if actor=='system' else 'resolved'
    doc['resolution']=f"{amount/100:,.2f} USD posted to {inv['id']}"
    audit(state,'remittance',doc['id'],'Posted',doc['resolution'],actor)
    notify(state,'remittance',doc['id'],'ledger',doc['resolution'])


def resolve(state,entity,action,data):
    book=state['remittance'];doc=next(d for d in book['documents'] if d['id']==entity)
    if doc['status']!='review': raise ValueError('Only pending reviews can be resolved')
    if action=='reject':
        doc['status']='rejected';doc['resolution']=data.get('note','Rejected after review')
        audit(state,'remittance',entity,'Rejected',doc['resolution'],'reviewer')
    elif action=='approve':
        inv=next((i for i in book['invoices'] if i['id']==data.get('invoice')),None)
        if not inv: raise ValueError('Select a valid invoice')
        ref=data.get('reference','').strip()
        if not ref: raise ValueError('A verified payment reference is required')
        if not data.get('note','').strip(): raise ValueError('Explain the evidence used for this correction')
        audit(state,'remittance',entity,'Correction approved',data['note'],'reviewer')
        post(state,doc,inv,ref,money(data.get('amount')),'reviewer')
    else: raise ValueError('Unknown action')
    return doc


def samples():
    def text(customer,invoice,amount,reference,currency='USD'):
        return f'Customer: {customer}\nInvoice: {invoice}\nAmount: {amount}\nReference: {reference}\nCurrency: {currency}'
    exact=text('Northline Studio','INV-1001','1250.00','BANK-001')
    return [('exact-payment.txt',exact),('partial-payment.txt',text('Cedar Works','INV-1002','300.00','BANK-002')),
            ('duplicate.txt',exact),('wrong-invoice.txt',text('Harbor Design','INV-9999','2100.00','BANK-004')),
            ('unknown-customer.txt',text('Unknown Demo Customer','INV-1004','640.00','BANK-005')),
            ('missing-invoice.txt',text('Cedar Works','','950.00','BANK-006')),
            ('overpayment.txt',text('Harbor Design','INV-1006','1900.00','BANK-007')),
            ('malformed.txt','Payment attached. No structured payment fields.'),
            ('missing-reference.txt',text('Northline Studio','INV-1007','720.00','')),
            ('exact-payment-2.txt',text('Cedar Works','INV-1008','320.00','BANK-010'))]
